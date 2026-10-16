import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L

df = pd.read_csv('grade_raw_T8_T10.csv', dtype=str).fillna('')
df['region'] = df.team.map(lambda t: 'North' if t.startswith(('HN', 'North')) else 'South')
N = len(df); last = N + 1
F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F4E78')
SUB = PatternFill('solid', fgColor='DDEBF7'); B = Font(name=F, bold=True); R = Font(name=F)
thin = Side(style='thin', color='BFBFBF'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
GR = '{"A+","A","B+","B","B-","C"}'

wb = Workbook()
# ---------- Data ----------
d = wb.active; d.title = 'Data'
cols = ['No', 'Team', 'Region', 'Segment', 'Status T10 MTD', 'Status T9', 'MOB T8', 'MOB T9', 'MOB T10',
        'Grade T8', 'Grade T9', 'Grade T10 MTD']
d.append(cols)
for i, r in df.iterrows():
    d.append([i + 1, r.team, r.region, r.seg, r.stMTD, r.stM1, r.mobM2, r.mobM1, r.mobMTD, r.M2, r.M1, r.MTD])
for c in d[1]: c.font, c.fill, c.alignment = H, HF, Alignment(horizontal='center', wrap_text=True)
for row in d.iter_rows(min_row=2):
    for c in row: c.font = R
for i, w in enumerate([6, 11, 9, 9, 12, 11, 9, 9, 9, 10, 10, 13], 1): d.column_dimensions[L(i)].width = w
d.freeze_panes = 'A2'; d.auto_filter.ref = f'A1:L{last}'

rng = lambda col: f'Data!${col}$2:${col}${last}'
MONTHS = [('T8', 'J', 'G'), ('T9', 'K', 'H'), ('T10 MTD', 'L', 'I')]
DIM = {'Region': 'C', 'Segment': 'D', 'Team': 'B'}

def summary(ws, mob3):
    title = 'Tỷ lệ BD Grade B- — T8 / T9 / T10 MTD (15/10/2026)' + (' — chỉ BD MOB3+' if mob3 else ' — toàn bộ BD có grade')
    ws['A1'] = title; ws['A1'].font = Font(name=F, bold=True, size=14)
    ws['A2'] = '%B- = #B- ÷ #BD có grade hợp lệ (A+/A/B+/B/B-/C) trong tháng. Loại grade "-" (MOB0) và ô trống (chưa vào / đã nghỉ). Công thức COUNTIFS trỏ về sheet Data.'
    ws['A2'].font = Font(name=F, italic=True, size=9, color='595959')
    row = 4
    blocks = [('Total', None, ['Total']),
              ('Region', 'C', ['North', 'South']),
              ('Segment', 'D', ['FS', 'SME', 'BCI']),
              ('Region × Segment', 'RS', [('North', 'FS'), ('North', 'SME'), ('South', 'FS'), ('South', 'SME'), ('South', 'BCI')]),
              ('Team', 'B', sorted(df.team.unique(), key=lambda t: (0 if t.startswith(('HN', 'North')) else 1, t)))]
    pct_cells = []
    for name, col, groups in blocks:
        ws.cell(row, 1, name).font = Font(name=F, bold=True, size=12); row += 1
        hdr1 = [name] + sum([[m, '', ''] for m, _, _ in MONTHS], []) + ['Δ %B- T10 vs T8']
        hdr2 = [''] + ['#B-', '#Graded', '%B-'] * 3 + ['(điểm %)']
        for j, v in enumerate(hdr1, 1): ws.cell(row, j, v)
        for j, v in enumerate(hdr2, 1): ws.cell(row + 1, j, v)
        for k in range(3): ws.merge_cells(start_row=row, start_column=2 + 3 * k, end_row=row, end_column=4 + 3 * k)
        for rr in (row, row + 1):
            for j in range(1, 12):
                c = ws.cell(rr, j); c.font, c.fill, c.alignment, c.border = H, HF, Alignment(horizontal='center'), BOX
        row += 2; first = row
        for g in groups:
            label = g if isinstance(g, str) else f'{g[0]}-{g[1]}'
            ws.cell(row, 1, label)
            if col is None: crit = ''
            elif col == 'RS': crit = f',{rng("C")},"{g[0]}",{rng("D")},"{g[1]}"'
            else: crit = f',{rng(col)},"{g}"'
            for k, (m, gc, mc) in enumerate(MONTHS):
                mobc = f',{rng(mc)},"MOB3+"' if mob3 else ''
                cb, cn, cp = 2 + 3 * k, 3 + 3 * k, 4 + 3 * k
                ws.cell(row, cb, f'=COUNTIFS({rng(gc)},"B-"{crit}{mobc})')
                ws.cell(row, cn, f'=SUMPRODUCT(COUNTIFS({rng(gc)},{GR}{crit}{mobc}))')
                ws.cell(row, cp, f'=IF({L(cn)}{row}=0,"",{L(cb)}{row}/{L(cn)}{row})')
                ws.cell(row, cp).number_format = '0.0%'
            ws.cell(row, 11, f'=IF(OR(D{row}="",J{row}=""),"",(J{row}-D{row})*100)')
            ws.cell(row, 11).number_format = '+0.0;-0.0;0.0'
            for j in range(1, 12):
                c = ws.cell(row, j); c.font = B if j in (1,) else R; c.border = BOX
                if j > 1: c.alignment = Alignment(horizontal='center')
            if name == 'Total':
                for j in range(1, 12): ws.cell(row, j).fill = SUB; ws.cell(row, j).font = B
            row += 1
        for k in range(3):
            pct_cells.append(f'{L(4 + 3 * k)}{first}:{L(4 + 3 * k)}{row - 1}')
        row += 1
    for rgn in pct_cells:
        ws.conditional_formatting.add(rgn, ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                                                          end_type='num', end_value=0.3, end_color='F8696B'))
    ws.column_dimensions['A'].width = 16
    for j in range(2, 11): ws.column_dimensions[L(j)].width = 9.5
    ws.column_dimensions['K'].width = 16
    ws['A' + str(row)] = 'Lưu ý: team chỉ 10–14 BD → 1 BD ≈ 7–10 điểm %. T10 là MTD 15/10, chưa chốt tháng.'
    ws['A' + str(row)].font = Font(name=F, italic=True, size=9, color='C00000')
    ws.freeze_panes = 'B4'

summary(wb.create_sheet('Summary'), False)
summary(wb.create_sheet('Summary MOB3+'), True)

# ---------- Transitions ----------
t = wb.create_sheet('Transitions')
t['A1'] = 'Ma trận dịch chuyển grade (số BD) — hàng = tháng trước, cột = tháng sau'; t['A1'].font = Font(name=F, bold=True, size=14)
G = ['A+', 'A', 'B+', 'B', 'B-', 'C', '-', '(trống)']
crit = lambda c, v: f'{rng(c)},"{v}"' if v != '(trống)' else f'{rng(c)},""'
row = 3
for title, a, b in [('T8 → T9', 'J', 'K'), ('T9 → T10 MTD', 'K', 'L')]:
    t.cell(row, 1, title).font = Font(name=F, bold=True, size=12); row += 1
    t.cell(row, 1, 'Từ \\ Đến')
    for j, g in enumerate(G, 2): t.cell(row, j, g)
    t.cell(row, len(G) + 2, 'Tổng')
    for j in range(1, len(G) + 3):
        c = t.cell(row, j); c.font, c.fill, c.alignment, c.border = H, HF, Alignment(horizontal='center'), BOX
    row += 1; top = row
    for g in G:
        t.cell(row, 1, g)
        for j, h in enumerate(G, 2):
            t.cell(row, j, f'=COUNTIFS({crit(a, g)},{crit(b, h)})')
        t.cell(row, len(G) + 2, f'=SUM(B{row}:{L(len(G) + 1)}{row})')
        for j in range(1, len(G) + 3):
            c = t.cell(row, j); c.border = BOX; c.alignment = Alignment(horizontal='center')
            c.font = B if (j == 1 or g == 'B-') else R
            if g == 'B-' or (j >= 2 and j <= len(G) + 1 and G[j - 2] == 'B-'): c.fill = PatternFill('solid', fgColor='FFF2CC')
        row += 1
    br = top + G.index('B-')
    row += 1
    t.cell(row, 1, f'Số phận BD B- ({title.split(" → ")[0]})').font = B; row += 1
    items = [('Lên ≥B', f'=SUM(B{br}:E{br})'), ('Giữ B-', f'=F{br}'), ('Rớt C', f'=G{br}'),
             ('Không còn grade / nghỉ', f'=H{br}+I{br}'), ('Tổng B-', f'=J{br}')]
    s0 = row
    for lab, f in items:
        t.cell(row, 1, lab).font = R; t.cell(row, 2, f).font = R
        t.cell(row, 3, f'=IF($B${s0 + 4}=0,"",B{row}/$B${s0 + 4})'); t.cell(row, 3).number_format = '0%'
        row += 1
    row += 1
    col_b = 6  # B- column index in matrix
    t.cell(row, 1, f'Nguồn BD B- ({title.split(" → ")[1]})').font = B; row += 1
    s1 = row
    items = [('Rớt từ ≥B', f'=SUM(F{top}:F{top + 3})'), ('Giữ B-', f'=F{br}'), ('Lên từ C', f'=F{top + 5}'),
             ('BD mới / chưa grade', f'=F{top + 6}+F{top + 7}'), ('Tổng B-', f'=SUM(F{top}:F{top + 7})')]
    for lab, f in items:
        t.cell(row, 1, lab).font = R; t.cell(row, 2, f).font = R
        t.cell(row, 3, f'=IF($B${s1 + 4}=0,"",B{row}/$B${s1 + 4})'); t.cell(row, 3).number_format = '0%'
        row += 1
    row += 2
t.column_dimensions['A'].width = 24
for j in range(2, 12): t.column_dimensions[L(j)].width = 9

# ---------- Notes ----------
n = wb.create_sheet('Ghi chú')
notes = [
    ('Nguồn', 'Google Sheet "[Internal SC] Grade C monitoring" → sheet "[SC] Data raw C movement" (lấy 16/10/2026). M-2 = T8, M-1 = T9, MTD = T10 chốt 15/10/2026.'),
    ('Phạm vi', f'{N} BD có trong sheet nguồn (gồm cả BD đã nghỉ trong kỳ). Đã bỏ email BD.'),
    ('Định nghĩa %B-', '#BD grade B- ÷ #BD có grade A+/A/B+/B/B-/C trong tháng đó. Loại grade "-" (MOB0) và ô trống.'),
    ('Giả định – Region', 'Suy từ tên team: HN*/North* → North; HCM*/BCI HCM → South.'),
    ('Giả định – Team/Segment', 'Theo team MTD (hoặc team cuối cùng nếu đã nghỉ). BD chuyển team trong kỳ được tính vào team hiện tại.'),
    ('Giới hạn', 'T10 là MTD 15/10, chưa chốt. Team 10–14 BD → 1 BD ≈ 7–10 điểm %, đọc như tín hiệu, không phải kết luận.'),
    ('Cách cập nhật', 'Dán đè data mới vào sheet Data (giữ đúng cột). Các sheet Summary và Transitions tự tính lại bằng COUNTIFS. Nếu số dòng tăng, mở rộng range trong công thức (hiện tới dòng ' + str(last) + ').'),
]
n['A1'] = 'Ghi chú phương pháp'; n['A1'].font = Font(name=F, bold=True, size=14)
for i, (k, v) in enumerate(notes, 3):
    n.cell(i, 1, k).font = B; n.cell(i, 2, v).font = R; n.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')
n.column_dimensions['A'].width = 24; n.column_dimensions['B'].width = 110

wb.move_sheet('Summary', -2); wb.move_sheet('Summary MOB3+', -2)
wb._sheets = [wb['Summary'], wb['Summary MOB3+'], wb['Transitions'], wb['Data'], wb['Ghi chú']]
wb.active = 0
wb.save('BD_Grade_B-_T8-T10.xlsx')
